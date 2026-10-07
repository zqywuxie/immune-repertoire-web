"""Preview and apply scoped manual registrations using existing sample records."""
from __future__ import annotations

import csv
import io
from collections import Counter
from uuid import uuid4

from flask import current_app
from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer
from openpyxl import Workbook, load_workbook
from sqlalchemy import or_
from sqlalchemy.exc import SQLAlchemyError

from flask_app.exceptions import StorageError, ValidationError
from flask_app.models.database import Project, SampleRecord, db
from flask_app.services.project_input_summary import iter_input_sample_ids
from flask_app.services.user_scope import assert_owned, current_user_id

FIELD_LABELS = {
    'sample_name': '样本名称', 'sequence_id': '序列编号', 'spices': '物种',
    'institution': '所属机构', 'chain_flag': '链标记', 'is_healthy': '健康状态',
    'illness': '疾病', 'is_pe': '双端测序', 'contain_method': '纳入方法', 'iso_tag': '同型标签',
}
HEADERS = {'record_id': '登记记录标识', 'sample_id': '原始样本编号', **FIELD_LABELS, 'clear_fields': '清空字段'}
MAX_ROWS = 5000
MAX_FILE_BYTES = 10 * 1024 * 1024


def _identifier(record):
    return (record.extra_metadata or {}).get('input_sample_id') or record.sample_id


def _snapshot(record):
    metadata = record.extra_metadata or {}
    return {'id': record.id, 'sample_id': record.sample_id,
            'fields': {field: getattr(record, field) for field in FIELD_LABELS},
            'source': {key: metadata.get(key) for key in ('source_asset_id', 'asset_set', 'registration_kind',
                                                        'input_sample_id', 'manual_fields', 'unmatched_input')}}


def _serializer():
    return URLSafeTimedSerializer(current_app.secret_key, salt='sample-registration-preview-v1')


class SampleBatchService:
    def __init__(self, asset_service):
        self.assets = asset_service

    def _scope(self, project, dataset):
        assert_owned(project, 'Project')
        if not isinstance(dataset, str) or not dataset.strip() or len(dataset.strip()) > 120:
            raise ValidationError(message='请先明确本次登记的项目和数据集。')
        return dataset.strip()

    def _coverage(self, project, dataset):
        return set(iter_input_sample_ids(self.assets.asset_query(project.id, inputs_only=True, asset_set=dataset)))

    def _records(self, project, dataset, identifiers, record_ids=(), *, lock=False):
        query = SampleRecord.query.filter(SampleRecord.project_id == project.id,
                                         SampleRecord.extra_metadata['asset_set'].as_string() == dataset)
        if identifiers or record_ids:
            query = query.filter(or_(SampleRecord.sample_id.in_(identifiers),
                                     SampleRecord.extra_metadata['input_sample_id'].as_string().in_(identifiers),
                                     SampleRecord.id.in_(record_ids)))
        if lock:
            query = query.populate_existing().with_for_update()
        return query.all()

    def template(self, project, dataset, record_ids=None):
        dataset = self._scope(project, dataset)
        if record_ids is not None and (not isinstance(record_ids, list) or not record_ids
                                      or len(record_ids) > MAX_ROWS or any(not isinstance(value, str) for value in record_ids)):
            raise ValidationError(message='请选择有效的登记记录。')
        if record_ids:
            records = self._records(project, dataset, [], record_ids)
            if {record.id for record in records} != set(record_ids):
                raise ValidationError(message='所选登记不属于本次项目和数据集，或已被移除。')
            rows = [(_identifier(record), record) for record in records]
        else:
            known = self._coverage(project, dataset)
            records = self._records(project, dataset, [])
            grouped = {}
            for record in records:
                grouped.setdefault(_identifier(record), []).append(record)
            rows = [(identifier, record) for identifier in sorted(known | {value for value in grouped if value})
                    for record in (grouped.get(identifier) or [None])]
        if len(rows) > MAX_ROWS:
            raise ValidationError(message=f'此数据集超过 {MAX_ROWS} 条，请在登记列表选择需要维护的记录后下载模板。')
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = '样本登记'
        sheet.append(list(HEADERS.values()))
        for identifier, record in rows:
            values = {'record_id': record.id if record else '', 'sample_id': identifier or '',
                      **({field: getattr(record, field) or '' for field in FIELD_LABELS} if record else {}), 'clear_fields': ''}
            sheet.append([values.get(key, '') for key in HEADERS])
        for row in sheet:
            for cell in row:
                cell.number_format = '@'
                cell.data_type = 's'
        sheet.freeze_panes = 'C2'
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 22
        guide = workbook.create_sheet('填写说明')
        for message in [f'本模板仅用于项目：{project.name}；数据集：{dataset}。',
                        '原始样本编号与登记记录标识用于匹配，不可更改。编号按文本填写，保留 001 等前导零。',
                        '不修改的列可以删除；空白单元格默认保留现有值。',
                        '清空字段列填写中文字段名称，以顿号或逗号分隔，如：疾病、所属机构。样本名称不能清空。',
                        '保存前先预览逐行变更；预览后发生的修改会提示冲突，不覆盖其他人的更新。',
                        '未匹配编号只有明确允许补充登记时才会新增，不会加入分析输入或增加输入覆盖。']:
            guide.append([message])
        guide.column_dimensions['A'].width = 100
        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return buffer

    def read_file(self, uploaded):
        content = uploaded.stream.read(MAX_FILE_BYTES + 1)
        if not content or len(content) > MAX_FILE_BYTES:
            raise ValidationError(message='请上传不超过 10 MB 的登记表。')
        suffix = str(uploaded.filename or '').lower()
        try:
            if suffix.endswith('.xlsx'):
                workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
                sheet = workbook['样本登记'] if '样本登记' in workbook.sheetnames else workbook.active
                values = iter(sheet.iter_rows(values_only=True))
                headers = list(next(values, []))
                raw_rows = []
                for row in values:
                    if len(raw_rows) >= MAX_ROWS:
                        if any(value not in (None, '') for value in row):
                            raise ValidationError(message=f'每批最多 {MAX_ROWS} 行，请分批上传。')
                        continue
                    if any(value not in (None, '') for value in row):
                        raw_rows.append(list(row))
                workbook.close()
            elif suffix.endswith('.csv'):
                values = csv.reader(io.StringIO(content.decode('utf-8-sig')), strict=True)
                headers = list(next(values, []))
                raw_rows = []
                for row in values:
                    if any(str(value).strip() for value in row):
                        raw_rows.append(row)
                        if len(raw_rows) > MAX_ROWS:
                            raise ValidationError(message=f'每批最多 {MAX_ROWS} 行，请分批上传。')
            else:
                raise ValidationError(message='登记表支持 Excel（.xlsx）或 UTF-8 CSV 文件。')
        except ValidationError:
            raise
        except Exception as error:
            raise ValidationError(message='登记表无法读取，请使用下载的模板或 UTF-8 CSV。') from error
        aliases = {**{key: key for key in HEADERS}, **{label: key for key, label in HEADERS.items()}}
        columns = [aliases.get(str(value or '').strip()) for value in headers]
        if not headers or any(column is None for column in columns) or len(set(columns)) != len(columns) or 'sample_id' not in columns:
            raise ValidationError(message='请保留原始样本编号列；列名不能重复或包含不支持的字段。')
        rows = []
        for index, raw in enumerate(raw_rows):
            row = dict(zip(columns, raw))
            if len(raw) > len(columns) and any(value not in (None, '') for value in raw[len(columns):]):
                row['_error'] = '此行列数超出表头，请检查分隔符。'
            identifier = row.get('sample_id')
            if identifier is not None and not isinstance(identifier, str):
                row['_error'] = '原始样本编号必须按文本填写，数字单元格可能丢失前导零。'
            row['fields'] = {key: value for key, value in row.items() if key in FIELD_LABELS}
            rows.append({key: value for key, value in row.items() if key not in FIELD_LABELS})
        return rows

    def preview(self, project, dataset, rows, allow_unmatched=False, record_ids=None):
        dataset = self._scope(project, dataset)
        if not isinstance(rows, list) or not rows or len(rows) > MAX_ROWS or not isinstance(allow_unmatched, bool):
            raise ValidationError(message=f'请提供 1 至 {MAX_ROWS} 行登记变更。')
        if any(not isinstance(row, dict) for row in rows):
            raise ValidationError(message='登记行格式无效。')
        selected_ids = None
        if record_ids is not None:
            if (not isinstance(record_ids, list) or not record_ids or len(record_ids) > MAX_ROWS
                    or any(not isinstance(value, str) or not value or len(value) > 36 for value in record_ids)):
                raise ValidationError(message='请选择有效的登记记录范围。')
            selected_ids = set(record_ids)
            selected = self._records(project, dataset, [], record_ids)
            if {record.id for record in selected} != selected_ids:
                raise ValidationError(message='所选登记不属于本次项目和数据集，或已被移除。请重新选择。')
        known = self._coverage(project, dataset)
        identifiers = [str(row.get('sample_id') or '').strip() for row in rows]
        counts = Counter(('record', row['record_id']) if isinstance(row.get('record_id'), str) and row['record_id']
                         else ('sample', identifier) for row, identifier in zip(rows, identifiers))
        records = self._records(project, dataset, identifiers, [str(row.get('record_id') or '') for row in rows])
        by_id = {record.id: record for record in records}
        grouped = {}
        for record in records:
            grouped.setdefault(_identifier(record), []).append(record)
        plans, public = [], []
        clear_aliases = {**{key: key for key in FIELD_LABELS}, **{label: key for key, label in FIELD_LABELS.items()}}
        for number, (raw, identifier) in enumerate(zip(rows, identifiers), 1):
            status, message, record, changes = 'invalid', '', None, {}
            record_id = raw.get('record_id') or ''
            try:
                if raw.get('_error'):
                    raise ValidationError(message=str(raw['_error']))
                if set(raw) - {'record_id', 'sample_id', 'fields', 'clear_fields', '_error'}:
                    raise ValidationError(message='此行包含不支持的登记字段。')
                if not isinstance(raw.get('sample_id'), str) or not identifier or len(identifier) > 255:
                    raise ValidationError(message='请按文本填写有效的原始样本编号。')
                if record_id and (not isinstance(record_id, str) or len(record_id) > 36):
                    raise ValidationError(message='登记记录标识无效。')
                if selected_ids is not None and record_id not in selected_ids:
                    raise ValidationError(message='此行不在本次所选登记范围内，请保留模板中的登记记录标识。')
                identity = ('record', record_id) if record_id else ('sample', identifier)
                if counts[identity] > 1:
                    raise ValidationError(message='本批原始编号重复，请先确认重复行，不会自动合并。')
                if record_id:
                    record = by_id.get(record_id)
                    if record is None or _identifier(record) != identifier:
                        raise ValidationError(message='登记记录与项目、数据集或原始编号不一致。')
                else:
                    candidates = grouped.get(identifier, [])
                    if len(candidates) > 1:
                        raise ValidationError(message='此编号对应多条登记，请通过记录标识选择具体记录。')
                    record = candidates[0] if candidates else None
                fields = raw.get('fields', {})
                if not isinstance(fields, dict) or set(fields) - FIELD_LABELS.keys():
                    raise ValidationError(message='请仅修改模板支持的补充字段。')
                clear = raw.get('clear_fields') or []
                if isinstance(clear, str):
                    import re
                    clear = [value.strip() for value in re.split('[,，、;；]', clear) if value.strip()]
                if not isinstance(clear, list) or any(not isinstance(value, str) or value not in clear_aliases for value in clear):
                    raise ValidationError(message='清空字段必须填写模板中的字段名称。')
                clear = {clear_aliases[value] for value in clear}
                if 'sample_name' in clear:
                    raise ValidationError(message='样本名称不能清空。')
                for field, value in fields.items():
                    if isinstance(value, (dict, list, bool)):
                        raise ValidationError(message=f'{FIELD_LABELS[field]}的格式无效。')
                    value = str(value).strip() if value is not None else ''
                    if field in clear and value:
                        raise ValidationError(message=f'{FIELD_LABELS[field]}同时填写新值和清空，请选择一种操作。')
                    if len(value) > (120 if field in ('is_healthy', 'is_pe') else 255):
                        raise ValidationError(message=f'{FIELD_LABELS[field]}内容过长。')
                    if value:
                        changes[field] = value
                changes.update({field: None for field in clear})
                if record:
                    changes = {field: value for field, value in changes.items() if value != getattr(record, field)}
                    status = 'update' if changes else 'unchanged'
                    message = '修改已有登记' if changes else '没有字段变化'
                elif identifier in known:
                    status, message = 'new', '新增已识别输入的登记'
                elif allow_unmatched:
                    status, message = 'new', '补充登记，尚未关联输入；不加入分析样本'
                else:
                    status, message = 'unmatched', '当前输入中未识别此编号；如需补录，请明确允许未关联登记后重新预览。'
            except ValidationError as error:
                message = error.message
            expected = _snapshot(record) if record else None
            plans.append({'row': number, 'sample_id': identifier, 'record_id': record.id if record else None,
                          'changes': changes, 'expected': expected, 'status': status, 'message': message,
                          'known': identifier in known})
            public.append({'row': number, 'sample_id': identifier, 'record_id': record.id if record else None,
                           'status': status, 'message': message, 'linked': identifier in known,
                           'can_apply': status in ('new', 'update', 'unchanged'),
                           'changes': [{'field': field, 'label': FIELD_LABELS[field],
                                        'before': getattr(record, field) if record else None, 'after': value}
                                       for field, value in changes.items()]})
        payload = {'project_id': project.id, 'asset_set': dataset, 'user_id': current_user_id(),
                   'operation_id': str(uuid4()), 'rows': plans}
        return {'preview_token': _serializer().dumps(payload), 'asset_set': dataset, 'project_name': project.name,
                'rows': public, 'counts': dict(Counter(row['status'] for row in public))}

    def apply(self, project, token, row_numbers=None):
        assert_owned(project, 'Project')
        try:
            payload = _serializer().loads(token, max_age=24 * 60 * 60)
        except (BadSignature, SignatureExpired, TypeError) as error:
            raise ValidationError(message='预览已失效或内容无效，请重新预览；不会自动覆盖原登记。') from error
        if payload['project_id'] != project.id or payload['user_id'] != current_user_id():
            raise ValidationError(message='预览不属于当前账号和项目，请重新预览。')
        dataset = self._scope(project, payload['asset_set'])
        plans = payload['rows']
        if row_numbers is not None:
            if (not isinstance(row_numbers, list) or not row_numbers
                    or any(type(value) is not int for value in row_numbers)
                    or len(set(row_numbers)) != len(row_numbers)
                    or not set(row_numbers).issubset({row['row'] for row in plans})):
                raise ValidationError(message='请选择预览中的有效行进行保存或重试。')
            plans = [row for row in plans if row['row'] in row_numbers]
        # Serialize new registrations within this project; each existing row is also locked and reloaded.
        db.session.query(Project).filter_by(id=project.id).with_for_update().one()
        known = self._coverage(project, dataset)
        results = []
        for plan in plans:
            outcome = {'row': plan['row'], 'sample_id': plan['sample_id'], 'status': 'failed', 'message': plan['message']}
            if plan['status'] not in ('new', 'update', 'unchanged'):
                results.append(outcome)
                continue
            try:
                with db.session.begin_nested():
                    matches = self._records(project, dataset, [plan['sample_id']], [plan['record_id']] if plan['record_id'] else [], lock=True)
                    if plan['record_id']:
                        record = next((row for row in matches if row.id == plan['record_id']), None)
                    else:
                        if len(matches) > 1:
                            raise ValidationError(message='此编号已有多条登记，请重新核对。')
                        record = matches[0] if matches else None
                    receipt = (record.extra_metadata or {}).get('registration_receipt') if record else None
                    if receipt and receipt.get('operation_id') == payload['operation_id'] and receipt.get('row') == plan['row']:
                        outcome.update(status='saved', record_id=record.id, replayed=True,
                                       message='此行此前已保存，本次未重复写入或覆盖后续修改。')
                        results.append(outcome)
                        continue
                    if (_snapshot(record) if record else None) != plan['expected']:
                        raise ValidationError(message='预览后此登记已被修改、新增或移除，请重新预览；本次未覆盖。')
                    if (plan['sample_id'] in known) != plan['known']:
                        raise ValidationError(message='预览后输入样本范围已改变，请重新预览。')
                    if plan['status'] == 'unchanged':
                        outcome.update(status='unchanged', record_id=record.id, message='没有字段变化，未重新写入。')
                    else:
                        if record is None:
                            metadata = {'asset_set': dataset, 'registration_kind': 'manual'}
                            if plan['known']:
                                metadata['input_sample_id'] = plan['sample_id']
                            else:
                                metadata['unmatched_input'] = True
                            record = SampleRecord(project_id=project.id, sample_id=plan['sample_id'],
                                                  sample_name=plan['sample_id'], extra_metadata=metadata)
                            db.session.add(record)
                        for field, value in plan['changes'].items():
                            setattr(record, field, value)
                        metadata = dict(record.extra_metadata or {})
                        metadata['manual_fields'] = sorted(set(metadata.get('manual_fields') or []) | set(plan['changes']))
                        metadata['registration_receipt'] = {'operation_id': payload['operation_id'], 'row': plan['row']}
                        record.extra_metadata = metadata
                        db.session.flush()
                        outcome.update(status='saved', record_id=record.id, message='登记已保存。', replayed=False)
            except ValidationError as error:
                outcome['message'] = error.message
            except SQLAlchemyError:
                current_app.logger.exception('Sample registration row failed')
                outcome['message'] = '此行保存失败，请重试；其他成功行会保留。'
            results.append(outcome)
        try:
            db.session.commit()
        except SQLAlchemyError as error:
            db.session.rollback()
            raise StorageError(message='本批提交失败，请使用原预览重试以核对保存结果。') from error
        return {'rows': results, 'counts': dict(Counter(row['status'] for row in results)), 'asset_set': dataset}

"""Bounded sample registry exports using the same scoped query as the list."""
import csv
import io
import json
from pathlib import Path
import tempfile

from flask_app.models.database import Project, SampleRecord

FIELDS = [
    ('project_id', '项目标识'), ('project_name', '项目'),
    ('sample_id', '样本编号'), ('sample_name', '样本名称'),
    ('sequence_id', '序列编号'), ('spices', '物种'), ('institution', '所属机构'),
    ('chain_flag', '链标记'), ('is_healthy', '健康状态'), ('illness', '疾病'),
    ('is_pe', '双端测序'), ('contain_method', '纳入方法'), ('iso_tag', '同型标签'),
]
BUSINESS = [(key, label) for key, label in FIELDS if key != 'project_id'] + [
    ('asset_set', '数据集'), ('registration_source', '登记来源')]
TECHNICAL = [('id', '登记记录标识'), ('project_id', '项目标识'),
    ('source_asset_id', '来源文件标识'), ('input_sample_id', '输入样本编号'),
    ('extra_metadata_json', '补充信息（原始记录）')]

def _business_value(key, value):
    if key == 'spices' and value:
        aliases = {'human':'人', 'homo sapiens':'人', '人类':'人',
                   'mouse':'小鼠', 'mus musculus':'小鼠', 'other':'其他'}
        return aliases.get(str(value).strip().lower(), value)
    if key == 'is_healthy':
        return {'yes':'健康','no':'非健康'}.get(value, value)
    if key == 'is_pe':
        return {'yes':'是','no':'否'}.get(value, value)
    return value

def _source(metadata):
    if metadata.get('source_asset_id'): return '来源表导入'
    if metadata.get('registration_kind') == 'manual':
        return '人工补充（尚未关联输入）' if metadata.get('unmatched_input') else '人工补充'
    return '来源未记录'

def build_sample_export(query, *, format='csv', columns='legacy'):
    """Generate on disk before sending so a generation error remains an HTTP error.

    No whole-row list, DataFrame, or in-memory output copy is created. The caller
    hands the open file to send_file; the WSGI file wrapper closes it on completion.
    """
    query = query.order_by(SampleRecord.id.desc())
    if columns == 'legacy':
        # Existing machine exports include arbitrary metadata columns. Discover
        # those keys with a narrow cursor, without retaining records or values.
        extras = set()
        for (metadata,) in query.with_entities(SampleRecord.extra_metadata).yield_per(500):
            extras.update((metadata or {}).keys())
        keys = [key for key, _ in FIELDS]
        keys += [key for key in sorted(extras) if key not in keys]
        headers = keys
    else:
        fields = BUSINESS + (TECHNICAL if columns == 'technical' else [])
        keys = [key for key, _ in fields]
        headers = [label for _, label in fields]
    output = tempfile.TemporaryFile(mode='w+b')
    text = None
    workbook = None
    try:
        if format == 'xlsx':
            from openpyxl import Workbook
            from openpyxl.cell import WriteOnlyCell
            from openpyxl.styles import Font, PatternFill
            workbook = Workbook(write_only=True)
            sheet = workbook.create_sheet('样本登记')
            sheet.freeze_panes = 'A2'
            from openpyxl.utils import get_column_letter
            for index in range(1, len(headers)+1):
                sheet.column_dimensions[get_column_letter(index)].width = 22
            def append(values, header=False):
                cells=[]
                for value in values:
                    cell=WriteOnlyCell(sheet, value='' if value is None else str(value))
                    cell.data_type='s'
                    cell.number_format='@'
                    if header:
                        cell.font=Font(bold=True, color='FFFFFF')
                        cell.fill=PatternFill('solid', fgColor='2563EB')
                    cells.append(cell)
                sheet.append(cells)
            append(headers, True)
        else:
            text = io.TextIOWrapper(output, encoding='utf-8-sig', newline='')
            writer = csv.writer(text)
            writer.writerow(headers)
            append = writer.writerow
        count = 0
        for sample, project_name in query.with_entities(SampleRecord, Project.name).yield_per(500):
            metadata = sample.extra_metadata or {}
            values = {key: getattr(sample, key, None) for key, _ in FIELDS}
            values['project_name'] = project_name
            if columns == 'legacy':
                values.update(metadata)
            else:
                values = {key: _business_value(key, value) for key, value in values.items()}
                values.update(id=sample.id, asset_set=metadata.get('asset_set'),
                    source_asset_id=metadata.get('source_asset_id'), input_sample_id=metadata.get('input_sample_id'),
                    registration_source=_source(metadata), extra_metadata_json=json.dumps(metadata, ensure_ascii=False))
            append([values.get(key) for key in keys])
            count += 1
        if workbook:
            workbook.save(output)
        if text:
            text.flush(); text.detach(); text = None
        output.seek(0)
        return output, count
    except BaseException:
        if text:
            text.close()
        output.close()
        raise
    finally:
        if workbook:
            # write-only worksheets use their own XML temporary file. A failed
            # export never reaches the normal save() cleanup in openpyxl 3.1.
            for worksheet in workbook.worksheets:
                if not worksheet.closed:
                    worksheet.close()
                writer = worksheet._writer
                if writer and Path(writer.out).exists():
                    writer.cleanup()
            workbook.close()

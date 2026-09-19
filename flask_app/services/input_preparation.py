"""Explicit table preparation; originals stay unchanged and outputs are registered assets."""
from pathlib import Path
from uuid import uuid4
from flask_app.exceptions import ValidationError
from flask_app.models.database import ProjectAsset, db
from flask_app.services.path_access_service import PathAccessService


def analysis_input_path(asset):
    mapping = (asset.metadata_json or {}).get('input_preparation')
    if not mapping:
        return asset.storage_path
    source = PathAccessService.validate_read_path(asset.storage_path)
    stat = source.stat()
    if stat.st_size != mapping['source_size'] or stat.st_mtime_ns != mapping['source_mtime_ns']:
        raise ValidationError(message='原始文件已改变，请重新确认工作表和列映射。')
    prepared = db.session.get(ProjectAsset, mapping['prepared_asset_id'])
    if not prepared or prepared.project_id != asset.project_id or prepared.asset_type != 'prepared_input':
        raise ValidationError(message='整理后的输入已删除，请重新整理数据。')
    path = PathAccessService.validate_read_path(prepared.storage_path)
    stat = path.stat()
    if stat.st_size != mapping['prepared_size'] or stat.st_mtime_ns != mapping['prepared_mtime_ns']:
        raise ValidationError(message='整理后的输入已改变，请重新整理数据。')
    return str(path)


def prepare_table(asset, options, projects_root):
    from flask_app.services.input_table_schema import inspect_table_schema
    from flask_app.services.input_quality import validate_analysis_inputs
    from flask_app.routes.api_script_hub._common import _robust_read_csv
    kind = {'datapoint':'profile','expression':'transcriptome','cibersort':'deconvolution'}.get(asset.asset_type, asset.asset_type)
    if kind not in {'profile','transcriptome','deconvolution'}:
        raise ValidationError(message='此类资产不支持工作表和列映射。')
    source = PathAccessService.validate_read_path(asset.storage_path)
    before = source.stat()
    schema = inspect_table_schema(source, options.get('sheet_name'))
    if schema['requires_sheet_selection']:
        raise ValidationError(message='文件包含多个工作表，请明确选择数据工作表。')
    columns = schema['columns']
    if any(not column.strip() for column in columns) or len(set(columns)) != len(columns):
        raise ValidationError(message='表头包含空列名或重复列名，请先修正原始文件。')
    identifier = options.get('identifier_column')
    if not isinstance(identifier,str) or identifier not in columns:
        raise ValidationError(message='请选择样本编号列或基因编号列。')
    canonical = 'Gene' if kind == 'transcriptome' else 'sample'
    if canonical in columns and identifier != canonical:
        raise ValidationError(message='目标标准列名已存在，请核对编号列。')
    destination = (Path(projects_root) / str(asset.project_id) / 'prepared_inputs' / (uuid4().hex + '.csv')).resolve()
    if not destination.is_relative_to(Path(projects_root).resolve()):
        raise ValidationError(message='项目输入目录无效。')
    destination.parent.mkdir(parents=True,exist_ok=True)
    try:
        write_prepared_table(source, destination, schema, identifier, canonical)
        validate_analysis_inputs([{'asset_type':kind,'path':str(destination)}])
        after = source.stat()
        if (before.st_size,before.st_mtime_ns) != (after.st_size,after.st_mtime_ns):
            raise ValidationError(message='整理期间原始文件发生变化，请重试。')
        stat = destination.stat()
        prepared = ProjectAsset(project_id=asset.project_id,asset_type='prepared_input',original_name=destination.name,
            storage_path=str(destination),size=stat.st_size,metadata_json={'source_asset_id':asset.id,'source_kind':kind,
            'sheet_name':schema['selected_sheet'],'identifier_column':identifier})
        db.session.add(prepared);db.session.flush()
        mapping={'prepared_asset_id':prepared.id,'source_size':before.st_size,'source_mtime_ns':before.st_mtime_ns,
                 'prepared_size':stat.st_size,'prepared_mtime_ns':stat.st_mtime_ns,
                 'sheet_name':schema['selected_sheet'],'identifier_column':identifier}
        asset.metadata_json={**(asset.metadata_json or {}),'input_preparation':mapping}
        db.session.commit()
        return mapping
    except Exception:
        db.session.rollback()
        destination.unlink(missing_ok=True)
        raise


def write_prepared_table(source, destination, schema, identifier, canonical):
    """Bound text-table batches by width; retry decoding from an empty output."""
    import pandas as pd
    from flask_app.routes.api_script_hub._common import _CSV_ENCODINGS, _robust_read_csv

    def write_frames(frames):
        with destination.open('w', encoding='utf-8', newline='') as output:
            first = True
            for frame in frames:
                frame = frame.rename(columns={identifier: canonical})
                if canonical == 'Gene':
                    frame = frame[[canonical] + [name for name in frame.columns if name != canonical]]
                frame.to_csv(output, index=False, header=first)
                first = False

    options = {'dtype': str, 'keep_default_na': False}
    if source.suffix.lower() in {'.xlsx', '.xlsm'}:
        write_frames(iter_workbook_frames(source, schema))
        return
    if source.suffix.lower() == '.xls':
        # The legacy binary format retains its existing reader.
        options['sheet_name'] = schema['selected_sheet']
        write_frames([_robust_read_csv(source, **options)])
        return
    options['sep'] = '\t' if source.suffix.lower() == '.tsv' else ','
    batch_rows = max(1, min(5000, 250000 // max(1, len(schema['columns']))))
    for encoding in dict.fromkeys([*_CSV_ENCODINGS, 'latin-1']):
        try:
            with pd.read_csv(source, encoding=encoding, chunksize=batch_rows, **options) as reader:
                write_frames(reader)
            return
        except UnicodeError:
            continue
    raise ValidationError(message='无法解析输入文件编码，请转换为中文兼容的文本表格后重试。')


def iter_workbook_frames(source, schema):
    """Read the selected worksheet with bounded row batches and cached formula values."""
    import pandas as pd
    from openpyxl import load_workbook

    columns = schema['columns']
    batch_rows = max(1, min(5000, 250000 // max(1, len(columns))))
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook[schema['selected_sheet']]
        sheet.reset_dimensions()
        rows = sheet.iter_rows()
        next(rows, None)
        batch, pending_empty = [], 0
        for cells in rows:
            values = []
            for cell in cells:
                value = cell.value
                if value is None or cell.data_type == 'e':
                    value = ''
                elif cell.data_type == 'n' and isinstance(value, float) and value.is_integer():
                    value = int(value)
                values.append(str(value))
            if any(value != '' for value in values[len(columns):]):
                raise ValidationError(message='工作表数据超出表头列数，请补齐列名后重试。')
            values = values[:len(columns)] + [''] * max(0, len(columns) - len(values))
            if not any(values):
                pending_empty += 1
                continue
            # Retain interior empty rows, but omit worksheet formatting after the data.
            while pending_empty:
                batch.append([''] * len(columns))
                pending_empty -= 1
                if len(batch) >= batch_rows:
                    yield pd.DataFrame(batch, columns=columns)
                    batch = []
            batch.append(values)
            if len(batch) >= batch_rows:
                yield pd.DataFrame(batch, columns=columns)
                batch = []
        if batch:
            yield pd.DataFrame(batch, columns=columns)
        else:
            yield pd.DataFrame(columns=columns)
    finally:
        workbook.close()


def revalidate_prepared_sources(job):
    for ref in (job.get('payload') or {}).get('source_assets', []):
        asset = db.session.get(ProjectAsset, ref.get('asset_id'))
        if not asset:
            raise ValidationError(message='任务输入资产已删除，请重新选择数据。')
        current = analysis_input_path(asset)
        if Path(current).resolve() != Path(ref['path']).resolve():
            raise ValidationError(message='任务输入映射已改变，请重新提交分析。')


def reset_input_preparation(asset):
    metadata = dict(asset.metadata_json or {})
    previous = metadata.pop('input_preparation', None)
    if previous is not None:
        asset.metadata_json = metadata
        db.session.commit()
    # Registered prepared assets remain available for provenance and lifecycle cleanup.
    return previous is not None

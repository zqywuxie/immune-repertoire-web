import io
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
import pytest
from flask import Flask
from flask_login import LoginManager, login_user
from flask_app.models.database import db, User
from flask_app.services.path_access_service import PathAccessService
from flask_app.services.result_path_resolver import candidate_job_roots
from flask_app.exceptions import ValidationError

@pytest.fixture
def app(tmp_path):
    application=Flask(__name__)
    application.config.update(TESTING=True,REQUIRE_LOGIN=True,SECRET_KEY='test',SQLALCHEMY_DATABASE_URI='sqlite:///:memory:',SQLALCHEMY_TRACK_MODIFICATIONS=False,USER_DATA_ROOT=tmp_path/'users',RESULTS_FOLDER=tmp_path/'results',UPLOAD_FOLDER=tmp_path/'uploads')
    (tmp_path/'uploads').mkdir()
    db.init_app(application)
    manager=LoginManager(application)
    @manager.user_loader
    def user_loader(user_id):return db.session.get(User,int(user_id))
    with application.app_context():
        db.create_all()
        for number in [1,2]:
            home=tmp_path/'users'/str(number);home.mkdir(parents=True)
            db.session.add(User(id=number,username=f'user{number}',email=f'u{number}@example.org',password_hash='unused',home_path=str(home)))
        db.session.commit()
    yield application
    with application.app_context():db.session.remove();db.drop_all()

def test_paths_and_results_do_not_cross_accounts(app,tmp_path):
    allowed=tmp_path/'users'/'1'/'sample.csv';allowed.write_text('x\n1\n')
    private=tmp_path/'users'/'2'/'private.csv';private.write_text('x\n2\n')
    with app.test_request_context():
        login_user(db.session.get(User,1))
        assert PathAccessService.validate_read_path(allowed)==allowed.resolve()
        with pytest.raises(ValidationError):PathAccessService.validate_read_path(private)
        with pytest.raises(ValidationError):PathAccessService.validate_write_path(private)
        assert all(str(path).startswith(str((tmp_path/'results'/'1').resolve())) for path in candidate_job_roots(tmp_path/'results','script_hub','job1'))
        assert list(candidate_job_roots(tmp_path/'results','script_hub','../2/job'))==[]

def test_unowned_job_is_not_visible_to_signed_in_user(app):
    from flask_app.routes.api_jobs import _can_access_job
    with app.test_request_context():
        login_user(db.session.get(User,1))
        assert _can_access_job({'user_id':1})
        assert not _can_access_job({'user_id':2})
        assert not _can_access_job({'user_id':None})

def test_real_pdf_upload_and_image_extraction(app):
    import fitz
    from flask_app.routes.api.pdf_routes import bp
    app.register_blueprint(bp,url_prefix='/api')
    import matplotlib.pyplot as plt
    image=io.BytesIO();fig,ax=plt.subplots(figsize=(1,1));ax.plot([0,1]);fig.savefig(image,format='png');plt.close(fig)
    pdf=fitz.open();page=pdf.new_page();page.insert_image(fitz.Rect(20,20,120,120),stream=image.getvalue());content=pdf.tobytes();pdf.close()
    client=app.test_client()
    with client.session_transaction() as session:session['_user_id']='1';session['_fresh']=True
    uploaded=client.post('/api/pdf/upload',data={'file':(io.BytesIO(content),'report.pdf')})
    assert uploaded.status_code==201,uploaded.get_json()
    extracted=client.post('/api/pdf/extract-images',json={'file_ids':[uploaded.json['file_id']],'indices':[0]})
    assert extracted.status_code==200,extracted.get_json()
    assert extracted.json['success_count']==1
    assert extracted.json['extracted_images']['report.pdf'][0]['image']

def test_real_ppt_parse_and_session_isolation(app,tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    from flask_app.routes.api_ppt import ppt_bp,_ppt_session_dir
    app.register_blueprint(ppt_bp)
    presentation=Presentation();slide=presentation.slides.add_slide(presentation.slide_layouts[6]);slide.shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(1)).text='Sharing Analysis - TRA'
    import matplotlib.pyplot as plt
    picture=io.BytesIO();fig,ax=plt.subplots(figsize=(1,1));ax.plot([1,2]);fig.savefig(picture,format='png');plt.close(fig)
    picture.seek(0);slide.shapes.add_picture(picture,Inches(1),Inches(2),width=Inches(2))
    source=io.BytesIO();presentation.save(source);source.seek(0)
    client=app.test_client()
    with client.session_transaction() as session:session['_user_id']='1';session['_fresh']=True
    # Keep the test session under its isolated fixture directory.
    with patch('flask_app.routes.api_ppt.tempfile.gettempdir',return_value=str(tmp_path)):
        parsed=client.post('/api/ppt/analyze',data={'file':(source,'template.pptx')})
        assert parsed.status_code==200,parsed.get_json()
        assert parsed.json['slide_count']==1
        position=parsed.json['heatmap_slides'][0]['image_positions'][0]
        image_path=tmp_path/'users'/'1'/'replacement.png';image_path.write_bytes(picture.getvalue())
        replaced=client.post('/api/ppt/replace',json={'session_id':parsed.json['session_id'],'heatmaps':{'TRA':{position['metric']:str(image_path)}}})
        assert replaced.status_code==200,replaced.get_json()
        assert replaced.json['replaced_count']==1
        assert client.get(replaced.json['download_url']).status_code==200
        url=f"/api/ppt/download/{parsed.json['session_id']}/template.pptx"
        assert client.get(url).status_code==200
        with client.session_transaction() as session:session['_user_id']='2'
        assert client.get(url).status_code != 200

def test_storage_cleanup_keeps_recent_caches_and_data(tmp_path):
    import os,time
    from scripts.storage_maintenance import inventory
    (tmp_path/'flask_app').mkdir();(tmp_path/'frontend').mkdir()
    old=tmp_path/'flask_app'/'__pycache__';old.mkdir();(old/'old.pyc').write_bytes(b'cache')
    recent=tmp_path/'frontend'/'.vite';recent.mkdir();(recent/'recent').write_bytes(b'recent')
    data=tmp_path/'flask_app'/'data';data.mkdir();(data/'sample.csv').write_text('preserve')
    stamp=time.time()-10*86400;os.utime(old/'old.pyc',(stamp,stamp));os.utime(old,(stamp,stamp))
    report=inventory(tmp_path,clean=True)
    assert report['removed_bytes']==5 and not old.exists()
    assert recent.exists() and (data/'sample.csv').exists()


def test_backup_can_be_opened_and_preserves_records(tmp_path):
    import sqlite3
    from scripts.backup_sqlite import backup_database
    original=tmp_path/'original.db'
    with sqlite3.connect(original) as connection:
        connection.execute('CREATE TABLE sample (name TEXT)')
        connection.execute("INSERT INTO sample VALUES ('S1')")
    backup=backup_database(original,tmp_path/'backup.db')
    with sqlite3.connect(backup) as restored:
        assert restored.execute('SELECT name FROM sample').fetchall()==[('S1',)]


def test_storage_endpoint_counts_only_owned_uploads(app,tmp_path):
    from flask_app.routes.api.config_params import bp
    from flask_app.models.database import File
    app.register_blueprint(bp,url_prefix='/api')
    with app.app_context():
        db.session.add(File(id='one',name='one.csv',original_name='one.csv',storage_path=str(tmp_path/'one.csv'),size=10,user_id=1,mime_type='text/csv',columns=['sample'],row_count=1))
        db.session.add(File(id='two',name='two.csv',original_name='two.csv',storage_path=str(tmp_path/'two.csv'),size=50,user_id=2,mime_type='text/csv',columns=['sample'],row_count=1))
        db.session.commit()
    client=app.test_client()
    with client.session_transaction() as session:session['_user_id']='1';session['_fresh']=True
    response=client.get('/api/storage')
    assert response.status_code==200,response.get_json()
    assert response.json['files']==1 and response.json['file_bytes']==10

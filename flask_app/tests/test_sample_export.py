import csv
import io
import pytest
from unittest.mock import patch

from openpyxl import load_workbook
from sqlalchemy.orm import Query

from flask_app.models.database import db, Project, SampleRecord, User
from flask_app.services.sample_export_service import build_sample_export
from flask_app.services.sample_registry_service import get_sample_registry_service
from flask_app.tests.test_data_management_contracts import context

def seed(project, dataset='甲', **values):
    record=SampleRecord(project_id=project.id, sample_id='001', sample_name='名称',
                        extra_metadata={'asset_set':dataset,'source_asset_id':'旧文件','input_sample_id':'001'}, **values)
    db.session.add(record); db.session.commit(); return record

def test_scoped_chinese_exports_preserve_text_and_source(context):
    app,_,project=context
    seed(project, institution='甲机构', sequence_id='0007', spices='human', is_healthy='yes')
    seed(project, '乙', institution='乙机构')
    client=app.test_client(); scope={'project_id':project.id,'asset_set':'甲','format':'xlsx','columns':'technical'}
    response=client.get('/api/samples/export',query_string=scope)
    assert response.status_code==200 and response.headers['X-Export-Count']=='1'
    workbook=load_workbook(io.BytesIO(response.data),read_only=True)
    rows=list(workbook.active.rows); headers=[c.value for c in rows[0]]
    values={key:cell for key,cell in zip(headers,rows[1])}
    assert values['样本编号'].value=='001' and values['样本编号'].data_type=='s'
    assert values['序列编号'].value=='0007' and values['来源文件标识'].value=='旧文件'
    assert values['物种'].value=='人' and values['健康状态'].value=='健康'
    workbook.close();response.close()
    with client.get('/api/samples/export',query_string={**scope,'format':'csv','columns':'business'}) as response:
        rows=list(csv.DictReader(io.StringIO(response.data.decode('utf-8-sig'))))
        assert len(rows)==1 and rows[0]['样本编号']=='001' and rows[0]['所属机构']=='甲机构'
        assert '项目标识' not in rows[0] and rows[0]['登记来源']=='来源表导入'

def test_export_default_contract_and_batches_do_not_collect_whole_list(context):
    app,_,project=context
    db.session.add_all([SampleRecord(project_id=project.id,sample_id=f'{i:04d}',sample_name='目标',
        extra_metadata={'asset_set':'甲','自定义':'值'}) for i in range(1103)])
    db.session.commit()
    with patch.object(Query,'all',side_effect=AssertionError('export must not collect all records')):
        response=app.test_client().get('/api/samples/export',query_string={'project_id':project.id,'q':'目标','page':2,'page_size':50})
        rows=list(csv.DictReader(io.StringIO(response.data.decode('utf-8-sig'))))
        assert response.headers['X-Export-Count']=='1103'
        assert len(rows)==1103 and len({row['sample_id'] for row in rows})==1103
        assert rows[0]['自定义']=='值' and 'project_name' in rows[0] and 'asset_set' in rows[0]
        response.close()

def test_xlsx_formula_like_values_remain_text_and_empty_export_has_headers(context):
    app,_,project=context
    record=seed(project);record.sample_name='=1+1';db.session.commit()
    response=app.test_client().get('/api/samples/export',query_string={'project_id':project.id,'format':'xlsx','columns':'business'})
    workbook=load_workbook(io.BytesIO(response.data));cells=list(workbook.active.rows)
    name=cells[1][[cell.value for cell in cells[0]].index('样本名称')]
    assert name.value=='=1+1' and name.data_type=='s';workbook.close();response.close()
    response=app.test_client().get('/api/samples/export',query_string={'project_id':project.id,'asset_set':'不存在','columns':'business'})
    assert response.headers['X-Export-Count']=='0'
    assert response.data.decode('utf-8-sig').startswith('项目,样本编号,');response.close()
    assert app.test_client().get('/api/samples/export?format=pdf').status_code==400

def test_export_account_scope_and_temporary_file_lifetime(context):
    app,_,project=context
    owner=User(username='owner-export',email='owner-export@local',password_hash='synthetic')
    other=User(username='other-export',email='other-export@local',password_hash='synthetic')
    db.session.add_all([owner,other]);db.session.flush();project.user_id=owner.id
    foreign=Project(name='不可访问',user_id=other.id);db.session.add(foreign);db.session.commit()
    seed(project,institution='允许');seed(foreign,institution='不可访问')
    app.config.update(REQUIRE_LOGIN=True)
    client=app.test_client()
    with client.session_transaction() as session: session['_user_id']=str(owner.id);session['_fresh']=True
    opened=[]
    def build(*args,**kwargs):
        result=build_sample_export(*args,**kwargs);opened.append(result[0]);return result
    with patch('flask_app.services.sample_export_service.build_sample_export',side_effect=build):
        response=client.get('/api/samples/export')
        assert '允许' in response.data.decode('utf-8-sig') and '不可访问' not in response.data.decode('utf-8-sig')
        response.close();assert opened[0].closed

@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_generation_failure_closes_temporary_output(context, format):
    _,_,project=context;seed(project)
    import tempfile
    from openpyxl.worksheet._writer import ALL_TEMP_FILES
    original_temporary_files=set(ALL_TEMP_FILES)
    output=tempfile.TemporaryFile(mode='w+b')
    with patch('flask_app.services.sample_export_service.tempfile.TemporaryFile',return_value=output), patch(
        'flask_app.services.sample_export_service._business_value',side_effect=RuntimeError('synthetic failure')):
        with pytest.raises(RuntimeError,match='synthetic failure'):
            build_sample_export(get_sample_registry_service().sample_query(project_id=project.id),columns='business',format=format)
    assert output.closed
    assert set(ALL_TEMP_FILES)==original_temporary_files

import pytest
from flask_app.tests.test_profile_workflow import profile_app
from flask_app.services.input_quality import inspect_input_quality


def test_four_inputs_compare_full_sample_sets_and_preserve_identifiers(profile_app, tmp_path):
    profile = tmp_path / 'profile.csv'
    profile.write_text('sample,group,metric\n001,A,1\n002,B,2\n', encoding='utf-8')
    expression = tmp_path / 'expression.tsv'
    expression.write_text('Gene\t001\t003\nG1\t2\t3\n', encoding='utf-8')
    infiltration = tmp_path / 'infiltration.csv'
    infiltration.write_text('Mixture,T cells,B cells\n001,0.2,0.8\n002,0.1,0.9\n', encoding='utf-8')
    pep = tmp_path / '001__TRA.csv'; pep.write_text('CDR3(pep),V,J,copy\nAAA,TRAV1,TRAJ1,1\n', encoding='utf-8')
    quality = inspect_input_quality([str(pep)], str(profile), str(expression), str(infiltration))
    assert len(quality['inputs']) == 4
    assert quality['errors'] == []
    align = {item['kind']: item for item in quality['alignments']}
    assert align['transcriptome']['missing_samples'] == ['002']
    assert align['transcriptome']['extra_samples'] == ['003']
    assert align['deconvolution']['matched_count'] == 2
    assert align['pep']['missing_samples'] == ['002']


def test_duplicate_and_blank_sample_ids_are_reported_with_missing_fields(profile_app, tmp_path):
    path = tmp_path / 'profile.csv'
    path.write_text('sample,group,metric\n001,A,1\n001,,2\n,B,3\n', encoding='utf-8')
    quality = inspect_input_quality([], str(path), '', '')
    report = quality['inputs'][0]
    assert report['duplicate_samples'] == ['001']
    assert report['missing_sample_count'] == 1
    assert report['missing_fields'] == {'group': 1}
    assert report['status'] == 'invalid' and quality['errors']


def test_duplicate_expression_columns_not_hidden_by_pandas(profile_app, tmp_path):
    path = tmp_path / 'expression.csv'
    path.write_text('Gene,001,001\nG1,1,2\n', encoding='utf-8')
    report = inspect_input_quality([], '', str(path), '')
    assert report['inputs'][0]['duplicate_samples'] == ['001']
    assert report['errors']


def test_unrecognized_sample_column_requires_mapping(profile_app, tmp_path):
    path = tmp_path / 'input.csv'; path.write_text('未知编号,值\n001,1\n', encoding='utf-8')
    report = inspect_input_quality([], str(path), '', '')
    assert report['inputs'][0]['status'] == 'needs_mapping'
    assert report['warnings']


def test_inspection_uses_registered_dataset_not_client_paths(profile_app, tmp_path):
    from flask_app.models.database import Project, ProjectAsset, db
    project = Project(name='输入质量测试'); db.session.add(project); db.session.flush()
    for name, content in [('Set1', 'sample,group\nold,A\nold,B\n'), ('Set2', 'sample,group\n001,A\n002,B\n')]:
        path=tmp_path/f'{name}.csv';path.write_text(content,encoding='utf-8')
        db.session.add(ProjectAsset(project_id=project.id,asset_type='profile',original_name=path.name,storage_path=str(path),size=path.stat().st_size,metadata_json={'asset_set':name}))
    db.session.commit()
    response=profile_app.test_client().post('/api/script-hub/data-selection/inspect',json={'project_id':project.id,'asset_set':'Set2','profile_path':str(tmp_path/'Set1.csv')})
    assert response.status_code==400,response.json
    response=profile_app.test_client().post('/api/script-hub/data-selection/inspect',json={'project_id':project.id,'asset_set':'Set2','profile_path':str(tmp_path/'Set2.csv')})
    assert response.status_code==200,response.json
    assert response.json['input_quality']['errors']==[]
    assert response.json['input_quality']['inputs'][0]['sample_count']==2


def test_expression_resolution_stays_in_selected_project_dataset(profile_app, tmp_path):
    from flask_app.models.database import Project, ProjectAsset, db
    from flask_app.routes.api_script_hub._common import _transcriptome_path_from_request
    project = Project(name='转录组数据集隔离')
    db.session.add(project); db.session.flush()
    paths = {}
    for name in ('Set1', 'Set2'):
        path = tmp_path / f'{name}-expression.csv'
        path.write_text('Gene,001,002\nG1,1,2\n', encoding='utf-8')
        paths[name] = str(path)
        db.session.add(ProjectAsset(project_id=project.id, asset_type='transcriptome', original_name=path.name, storage_path=str(path), size=path.stat().st_size, metadata_json={'asset_set': name}))
    db.session.commit()
    selected = {'project_id': project.id, 'asset_set': 'Set2', 'expression_path': paths['Set1']}
    assert _transcriptome_path_from_request(selected, 'expression_path') == paths['Set2']
    selected['asset_set'] = 'MissingSet'
    assert _transcriptome_path_from_request(selected, 'expression_path') is None
    # Existing non-project tools can still resolve their explicit allowed path.
    assert _transcriptome_path_from_request({'expression_path': paths['Set1']}, 'expression_path') == paths['Set1']


def test_cache_submission_rechecks_changed_input_before_reuse(profile_app, tmp_path):
    import pytest
    from flask_app.exceptions import ValidationError
    from flask_app.routes.api_script_hub._common import _build_script_cache_context
    path = tmp_path / 'profile.csv'
    path.write_text('sample,group\n001,A\n002,B\n', encoding='utf-8')
    assets = [{'asset_type': 'profile', 'path': str(path)}]
    assert not inspect_input_quality([], str(path), '', '')['errors']
    path.write_text('sample,group\n001,A\n001,B\n', encoding='utf-8')
    with pytest.raises(ValidationError, match='输入数据检查未通过'):
        _build_script_cache_context(project_id=None, module_name='profile', input_paths=assets, config_json={})


def test_execution_quality_rechecks_current_content_and_skips_derived_tables(profile_app, tmp_path):
    import pytest
    from flask_app.exceptions import ValidationError
    from flask_app.services.input_quality import validate_analysis_inputs
    path = tmp_path / 'expression.csv'
    path.write_text('Gene,001,002\nG1,1,2\n', encoding='utf-8')
    inputs = [{'asset_type': 'transcriptome', 'path': str(path)}]
    validate_analysis_inputs(inputs)
    path.write_text('Gene,001,001\nG1,1,2\n', encoding='utf-8')
    with pytest.raises(ValidationError): validate_analysis_inputs(inputs)
    validate_analysis_inputs([{'asset_type': 'cached_usage', 'path': str(path)}])


def test_expression_numeric_scan_covers_late_rows_without_changing_values(profile_app, tmp_path):
    path = tmp_path / 'expression.tsv'
    path.write_text('Gene\t001\t002\n' + ''.join(f'G{i}\t-1.2\t3e-4\n' for i in range(2500)) + 'last\tNaN\t\n', encoding='utf-8')
    quality = inspect_input_quality([], '', str(path), '')
    numeric = quality['inputs'][0]['numeric_content']
    assert numeric['row_count'] == 2501 and numeric['invalid_count'] == 2
    assert numeric['examples'] == ['第 2502 行 / 001', '第 2502 行 / 002']
    assert quality['errors']


def test_infiltration_numeric_errors_and_empty_expression_are_rejected(profile_app, tmp_path):
    path = tmp_path / 'infiltration.csv'
    path.write_text('Mixture,T cells,B cells,Correlation\n001,abc,inf,-0.2\n', encoding='utf-8')
    quality = inspect_input_quality([], '', '', str(path))
    assert quality['inputs'][0]['numeric_content']['invalid_count'] == 2
    assert quality['errors']
    path.write_text('Gene,001,002\n', encoding='utf-8')
    assert inspect_input_quality([], '', str(path), '')['errors']


def test_workbook_schema_preserves_headers_and_selected_sheet(profile_app, tmp_path):
    import pytest
    from openpyxl import Workbook
    from flask_app.exceptions import ValidationError
    from flask_app.services.input_table_schema import inspect_table_schema
    path=tmp_path/'input.xlsx'
    workbook=Workbook();workbook.active.title='说明';workbook.active.append(['说明'])
    table=workbook.create_sheet('数据');table.append(['样本编码','001','001']);table.append(['0007',1,2])
    workbook.save(path)
    schema=inspect_table_schema(path)
    assert schema['sheets']==['说明','数据'] and schema['requires_sheet_selection']
    schema=inspect_table_schema(path,'数据')
    assert schema['columns']==['样本编码','001','001']
    assert schema['preview_rows'][0][0]=='0007'
    with pytest.raises(ValidationError):inspect_table_schema(path,'不存在')


def test_table_schema_endpoint_resolves_registered_dataset(profile_app, tmp_path):
    from flask_app.models.database import Project,ProjectAsset,db
    project=Project(name='结构检查');db.session.add(project);db.session.flush()
    path=tmp_path/'table.csv';path.write_text('样本编码,分组\n0007,甲\n',encoding='utf-8')
    db.session.add(ProjectAsset(project_id=project.id,asset_type='profile',storage_path=str(path),original_name=path.name,size=path.stat().st_size,metadata_json={'asset_set':'Set2'}));db.session.commit()
    client=profile_app.test_client()
    payload={'project_id':project.id,'asset_set':'Set2','kind':'profile','profile_path':'/unregistered.csv'}
    response=client.post('/api/script-hub/data-selection/table-schema',json=payload)
    assert response.status_code==400,response.json
    payload['profile_path']=str(path)
    response=client.post('/api/script-hub/data-selection/table-schema',json=payload)
    assert response.status_code==200,response.json
    assert response.json['preview_rows']==[['0007','甲']]
    payload['asset_set']='Missing'
    assert client.post('/api/script-hub/data-selection/table-schema',json=payload).status_code==400


def test_prepared_workbook_is_used_by_analysis_and_retains_original(profile_app, tmp_path):
    import pytest
    from openpyxl import Workbook
    from flask_app.exceptions import ValidationError
    from flask_app.models.database import Project, ProjectAsset, db
    from flask_app.services.input_preparation import prepare_table, analysis_input_path, revalidate_prepared_sources
    from flask_app.services.analysis_artifacts import capture_input_lineage
    from flask_app.routes.api_script_hub._common import _profile_path_from_request
    project=Project(name='工作表映射');db.session.add(project);db.session.flush()
    source=tmp_path/'input.xlsx';workbook=Workbook();workbook.active.title='说明';workbook.active.append(['说明'])
    sheet=workbook.create_sheet('数据');sheet.append(['编号','group','metric']);sheet.append(['0001','A',1]);sheet.append(['0002','B',2]);workbook.save(source)
    original=source.read_bytes()
    asset=ProjectAsset(project_id=project.id,asset_type='profile',storage_path=str(source),original_name=source.name,size=source.stat().st_size,metadata_json={'asset_set':'Set2'})
    db.session.add(asset);db.session.commit()
    from flask_app.routes.api_projects import project_api_bp
    profile_app.register_blueprint(project_api_bp)
    client=profile_app.test_client()
    response=client.post(f'/api/projects/{project.id}/assets/{asset.id}/prepare-input',json={'sheet_name':'数据','identifier_column':'编号'})
    assert response.status_code==200,response.json
    mapping=response.json['input_preparation']
    schema=client.post(f'/api/projects/{project.id}/assets/{asset.id}/input-schema',json={'sheet_name':'数据'})
    assert schema.status_code==200,schema.json
    assert schema.json['columns'][0]=='编号'
    prepared=analysis_input_path(asset)
    assert prepared != str(source) and source.read_bytes()==original
    assert _profile_path_from_request({'project_id':project.id,'asset_set':'Set2'},'profile_path')==prepared
    quality=inspect_input_quality([],prepared,'','')
    assert not quality['errors'] and quality['inputs'][0]['sample_count']==2
    lineage=capture_input_lineage(project.id,[{'path':prepared,'asset_type':'profile'}],'Set2')
    assert lineage['source_assets'][0]['asset_id']==asset.id
    assert db.session.get(ProjectAsset,mapping['prepared_asset_id']).asset_type=='prepared_input'
    revalidate_prepared_sources({'payload':lineage})
    source.write_bytes(original+b'changed')
    with pytest.raises(ValidationError,match='原始文件已改变'):revalidate_prepared_sources({'payload':lineage})


def test_reset_mapping_preserves_files_but_invalidates_queued_input(profile_app, tmp_path):
    import pytest
    from flask_app.exceptions import ValidationError
    from flask_app.models.database import Project,ProjectAsset,db
    from flask_app.services.input_preparation import prepare_table,analysis_input_path,revalidate_prepared_sources
    from flask_app.services.analysis_artifacts import capture_input_lineage
    from flask_app.routes.api_projects import project_api_bp
    profile_app.register_blueprint(project_api_bp)
    source=tmp_path/'source.csv';source.write_text('编号,group\n0001,A\n0002,B\n',encoding='utf-8')
    project=Project(name='重置映射');db.session.add(project);db.session.flush()
    asset=ProjectAsset(project_id=project.id,asset_type='profile',storage_path=str(source),original_name=source.name,size=source.stat().st_size)
    db.session.add(asset);db.session.commit()
    mapping=prepare_table(asset,{'identifier_column':'编号'},tmp_path/'projects')
    prepared=analysis_input_path(asset)
    lineage=capture_input_lineage(project.id,[{'path':prepared}])
    response=profile_app.test_client().delete(f'/api/projects/{project.id}/assets/{asset.id}/prepare-input')
    assert response.status_code==200 and response.json['changed']
    assert analysis_input_path(asset)==str(source)
    assert source.exists() and db.session.get(ProjectAsset,mapping['prepared_asset_id'])
    from pathlib import Path
    assert Path(prepared).exists()
    with pytest.raises(ValidationError,match='映射已改变'):revalidate_prepared_sources({'payload':lineage})


def test_multiple_registered_profiles_require_explicit_version(profile_app, tmp_path):
    from flask_app.models.database import Project, ProjectAsset, db
    from flask_app.routes.api_script_hub._common import _profile_path_from_request
    from flask_app.exceptions import ValidationError
    project=Project(name='多个版本');db.session.add(project);db.session.flush()
    for i in range(2):
        path=tmp_path/f'version{i}.csv';path.write_text(f'sample,value\nS1,{i}\n')
        db.session.add(ProjectAsset(project_id=project.id,asset_type='profile',storage_path=str(path),original_name=path.name,size=path.stat().st_size,metadata_json={'asset_set':'batch'}))
    db.session.commit()
    payload={'project_id':project.id,'asset_set':'batch'}
    with pytest.raises(ValidationError,match='多份'):
        _profile_path_from_request(payload,'profile_path')
    payload['profile_path']=str(tmp_path/'version1.csv')
    assert _profile_path_from_request(payload,'profile_path') == str(tmp_path/'version1.csv')


def test_joint_inputs_without_common_samples_are_rejected(profile_app,tmp_path):
    from flask_app.services.input_quality import validate_analysis_inputs
    from flask_app.exceptions import ValidationError
    profile=tmp_path/'profile.csv';profile.write_text('sample,value\nA,1\n')
    expression=tmp_path/'expression.csv';expression.write_text('gene,B\nG1,2\n')
    with pytest.raises(ValidationError,match='没有匹配'):
        validate_analysis_inputs([{'asset_type':'profile','path':str(profile)},{'asset_type':'transcriptome','path':str(expression)}])


def test_infiltration_groups_are_text_and_quality_columns_are_not_cells(profile_app, tmp_path):
    path = tmp_path/'infiltration.csv'
    content = 'Mixture,Group,T cells,B cells,P-value,Correlation,RMSE\n001,治疗前,0.2,0.8,0.01,0.9,0.2\n002,治疗后,0.4,0.6,0.03,0.8,0.3\n'
    path.write_text(content, encoding='utf-8')
    result = inspect_input_quality([], '', '', str(path))
    assert not result['errors']
    report = result['inputs'][0]
    assert report['samples'] == ['001','002']
    assert report['cell_columns'] == ['T cells','B cells']
    assert report['group_columns'] == ['Group']
    assert report['quality_columns'] == ['P-value','Correlation','RMSE']
    assert report['numeric_content']['invalid_count'] == 0
    assert path.read_text(encoding='utf-8') == content


def test_infiltration_quality_only_is_not_a_cell_matrix(profile_app, tmp_path):
    path = tmp_path/'quality.csv'
    path.write_text('Mixture,category,P-value,Correlation,RMSE\n001,A,0.1,0.9,0.2\n',encoding='utf-8')
    result = inspect_input_quality([], '', '', str(path))
    assert result['inputs'][0]['status'] == 'invalid'
    assert any('未包含细胞数值列' in message for message in result['errors'])


def test_infiltration_unknown_text_and_bad_numeric_are_not_silently_dropped(profile_app, tmp_path):
    path = tmp_path/'bad.csv'
    path.write_text('Mixture,分组,T cells,备注,P-value\n001,甲,NaN,未知字段,错误\n',encoding='utf-8')
    result = inspect_input_quality([], '', '', str(path))
    assert result['inputs'][0]['numeric_content']['invalid_count'] == 3
    assert result['errors']


def test_infiltration_workbook_keeps_group_text(profile_app, tmp_path):
    from openpyxl import Workbook
    path = tmp_path/'infiltration.xlsx'
    book = Workbook()
    book.active.append(['sample','分组','T cells','B cells','RMSE'])
    book.active.append(['001','甲',0.2,0.8,0.1])
    book.active.append(['002','乙',0.3,0.7,0.2])
    book.save(path)
    report = inspect_input_quality([], '', '', str(path))
    assert not report['errors']
    assert report['inputs'][0]['cell_columns'] == ['T cells','B cells']
    assert report['inputs'][0]['samples'] == ['001','002']


def test_infiltration_cell_columns_match_original_r_helper(profile_app, tmp_path):
    import os
    import json
    import subprocess
    from pathlib import Path
    reference = os.environ.get('REFERENCE_PIPELINE')
    if not reference:
        pytest.skip('需要只读挂载原始 pipeline 并设置 REFERENCE_PIPELINE')
    path = tmp_path/'infiltration.csv'
    path.write_text('Mixture,Group,T cells,B cells,P-value,Correlation,RMSE\nS1,A,0.2,0.8,0.01,0.9,0.2\nS2,B,0.3,0.7,0.02,0.8,0.3\n',encoding='utf-8')
    command = 'args <- commandArgs(TRUE); source(args[1]); x <- read_table(args[2]); cat(jsonlite::toJSON(resolve_columns(x, NULL, exclude=c("Mixture","Group","P-value","Correlation","RMSE"))))'
    result = subprocess.run(['Rscript','-e',command,str(Path(reference)/'07.immuneInfiltration'/'plot_deconv_helpers.R'),str(path)],check=True,capture_output=True,text=True)
    report = inspect_input_quality([], '', '', str(path))
    assert not report['errors']
    assert report['inputs'][0]['cell_columns'] == json.loads(result.stdout)

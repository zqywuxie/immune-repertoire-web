from pathlib import Path
import pandas as pd
import pytest
from flask_app.services.input_preparation import write_prepared_table


@pytest.mark.parametrize('suffix,sep,encoding', [('.csv', ',', 'utf-8'), ('.tsv', '\t', 'gb18030')])
def test_chunked_mapping_preserves_all_rows_and_values(tmp_path, monkeypatch, suffix, sep, encoding):
    source = tmp_path / ('input' + suffix)
    original = ('甲' + sep + '编号' + sep + '乙\n' + ''.join(
        f'{index}{sep}{index:07d}{sep}2.50\n' for index in range(12001)))
    source.write_text(original, encoding=encoding)
    destination = tmp_path / 'prepared.csv'
    read_csv = pd.read_csv
    batches = []
    def bounded_reader(*args, **kwargs):
        assert kwargs.get('chunksize', 0) > 0
        batches.append(kwargs['chunksize'])
        return read_csv(*args, **kwargs)
    monkeypatch.setattr(pd, 'read_csv', bounded_reader)
    write_prepared_table(source, destination, {'columns': ['甲','编号','乙'], 'selected_sheet':None}, '编号', 'Gene')
    result = read_csv(destination, dtype=str, keep_default_na=False)
    assert result.columns.tolist() == ['Gene', '甲', '乙']
    assert len(result) == 12001
    assert result.iloc[0].tolist() == ['0000000', '0', '2.50']
    assert result.iloc[-1].tolist() == ['0012000', '12000', '2.50']
    assert max(batches) <= 5000
    assert source.read_text(encoding=encoding) == original


def test_late_encoding_fallback_restarts_output(tmp_path):
    source = tmp_path / 'late.csv'
    rows = [f'{i:07d},' + 'a' * 120 for i in range(15000)]
    rows.append('0015000,中文')
    source.write_bytes(('identifier,value\n' + '\n'.join(rows) + '\n').encode('gb18030'))
    destination = tmp_path / 'prepared.csv'
    write_prepared_table(source, destination, {'columns':['identifier','value'], 'selected_sheet':None}, 'identifier', 'sample')
    result = pd.read_csv(destination, dtype=str, keep_default_na=False)
    assert len(result) == 15001
    assert result['sample'].is_unique
    assert result.iloc[-1].tolist() == ['0015000', '中文']


def test_workbook_batches_match_existing_reader(tmp_path, monkeypatch):
    from datetime import datetime
    from openpyxl import Workbook
    from openpyxl.styles import Font
    source = tmp_path / 'input.xlsx'
    book = Workbook()
    book.active.title = '说明'
    sheet = book.create_sheet('实验数据')
    sheet.append(['数值','编号','备注','时间'])
    for i in range(10003):
        sheet.append([i + 0.25, f'{i:07d}', '中文' if i % 2 else None, datetime(2026, 1, 1)])
    sheet.append([None, None, None, None])
    sheet.append([1.0, '0010004', True, None])
    sheet.cell(20000,1).font = Font(bold=True)
    book.save(source)
    original = source.read_bytes()
    expected = pd.read_excel(source, sheet_name='实验数据', dtype=str, keep_default_na=False)
    expected = expected.rename(columns={'编号':'Gene'})[['Gene','数值','备注','时间']]
    expected_path = tmp_path / 'expected.csv'
    expected.to_csv(expected_path, index=False)
    def reject_full_read(*args, **kwargs):
        raise AssertionError('workbook must not use a full pandas read')
    monkeypatch.setattr(pd, 'read_excel', reject_full_read)
    destination = tmp_path / 'prepared.csv'
    write_prepared_table(source, destination, {'columns':['数值','编号','备注','时间'], 'selected_sheet':'实验数据'}, '编号', 'Gene')
    assert destination.read_bytes() == expected_path.read_bytes()
    assert source.read_bytes() == original

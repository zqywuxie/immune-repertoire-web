"""Batch registration contracts using isolated records and cached synthetic coverage."""
import io

import pytest
from flask import Flask, g
from flask_login import LoginManager
from openpyxl import Workbook, load_workbook

from flask_app.exceptions import AppException
from flask_app.models.database import Project, ProjectAsset, SampleRecord, User, db
from flask_app.routes import api_projects
from flask_app.services.project_asset_service import ProjectAssetService
from flask_app.services.project_service import ProjectService


@pytest.fixture
def batch_context(tmp_path, monkeypatch):
    app = Flask(__name__)
    app.config.update(TESTING=True, REQUIRE_LOGIN=False, SECRET_KEY='batch-synthetic',
                      SQLALCHEMY_DATABASE_URI='sqlite:///:memory:', SQLALCHEMY_TRACK_MODIFICATIONS=False)
    db.init_app(app)
    login = LoginManager(app)
    login.user_loader(lambda identifier: db.session.get(User, int(identifier)))
    assets = ProjectAssetService(tmp_path / 'projects')
    monkeypatch.setattr(api_projects, '_asset_service', lambda: assets)
    monkeypatch.setattr(api_projects, '_project_service', lambda: ProjectService(tmp_path / 'projects'))
    app.register_blueprint(api_projects.project_api_bp)
    app.register_error_handler(AppException, lambda error: (error.to_dict(), error.http_status))
    with app.app_context():
        db.create_all()
        project = Project(name='登记试验')
        db.session.add(project)
        db.session.flush()
        for dataset in ('甲', '乙'):
            db.session.add(ProjectAsset(project_id=project.id, asset_type='profile', original_name='指标.csv',
                storage_path=f'/synthetic/{dataset}.csv', metadata_json={'asset_set': dataset,
                    'validation': {'status': 'valid', 'summary': {'inputs': [{'kind': 'profile', 'samples': ['001', '002', '003']}]}}}))
        db.session.commit()
        yield app, project
        db.session.remove()
        db.drop_all()


def record(project, identifier='001', dataset='甲', **fields):
    value = SampleRecord(project_id=project.id, sample_id=identifier, sample_name=identifier,
                         extra_metadata={'asset_set': dataset, 'source_asset_id': 'source-a'}, **fields)
    db.session.add(value)
    db.session.commit()
    return value


def preview(client, project, rows, dataset='甲', **extra):
    return client.post(f'/api/projects/{project.id}/samples/batch/preview',
                       json={'asset_set': dataset, 'rows': rows, **extra})


def apply(client, project, response, **extra):
    return client.post(f'/api/projects/{project.id}/samples/batch/apply',
                       json={'preview_token': response.json['preview_token'], **extra})


def test_preview_is_readonly_and_unmatched_requires_explicit_option(batch_context):
    app, project = batch_context
    with app.test_client() as client:
        response = preview(client, project, [{'sample_id': '001', 'fields': {'illness': '疾病'}}, {'sample_id': '999'}])
        assert response.status_code == 200
        assert [row['status'] for row in response.json['rows']] == ['new', 'unmatched']
        assert SampleRecord.query.count() == 0
        saved = apply(client, project, response)
        assert saved.json['counts'] == {'saved': 1, 'failed': 1}
        assert SampleRecord.query.one().sample_id == '001'


def test_blank_preserves_explicit_clear_and_manual_source_fields(batch_context):
    app, project = batch_context
    sample = record(project, illness='旧疾病', institution='旧机构')
    sample.extra_metadata = {**sample.extra_metadata, 'manual_fields': ['illness']}
    db.session.commit()
    with app.test_client() as client:
        response = preview(client, project, [{'record_id': sample.id, 'sample_id': '001',
            'fields': {'illness': '', 'sample_name': '新名称'}, 'clear_fields': ['所属机构']}])
        assert {change['field'] for change in response.json['rows'][0]['changes']} == {'sample_name', 'institution'}
        assert apply(client, project, response).json['counts'] == {'saved': 1}
        assert sample.illness == '旧疾病' and sample.institution is None and sample.sample_name == '新名称'
        assert sample.extra_metadata['source_asset_id'] == 'source-a'
        assert set(sample.extra_metadata['manual_fields']) == {'sample_name', 'institution', 'illness'}


def test_repeat_apply_and_intervening_edit_do_not_rewrite_success(batch_context):
    app, project = batch_context
    with app.test_client() as client:
        response = preview(client, project, [{'sample_id': '001', 'fields': {'illness': '首次'}}])
        assert apply(client, project, response).json['rows'][0]['replayed'] is False
        sample = SampleRecord.query.one()
        sample.illness = '后续修改'
        db.session.commit()
        result = apply(client, project, response)
        assert result.json['rows'][0]['replayed'] is True
        assert SampleRecord.query.count() == 1 and sample.illness == '后续修改'


def test_partial_conflict_retry_and_fresh_preview(batch_context):
    app, project = batch_context
    sample = record(project, illness='旧')
    with app.test_client() as client:
        response = preview(client, project, [{'sample_id': '001', 'fields': {'illness': '待保存'}}, {'sample_id': '002'}])
        sample.illness = '其他人修改'
        db.session.commit()
        result = apply(client, project, response)
        assert result.json['counts'] == {'failed': 1, 'saved': 1}
        assert sample.illness == '其他人修改'
        retried = apply(client, project, response, rows=[1])
        assert retried.json['counts'] == {'failed': 1}
        assert SampleRecord.query.count() == 2
        fresh = preview(client, project, [{'sample_id': '001', 'fields': {'illness': '确认修改'}}])
        assert apply(client, project, fresh).json['counts'] == {'saved': 1}
        assert sample.illness == '确认修改'


def test_unmatched_supplement_and_datasets_do_not_inflate_input_coverage(batch_context):
    app, project = batch_context
    with app.test_client() as client:
        response = preview(client, project, [{'sample_id': '999', 'fields': {'sample_name': '补录'}}], allow_unmatched=True)
        assert response.json['rows'][0]['linked'] is False
        assert apply(client, project, response).json['counts'] == {'saved': 1}
        sample = SampleRecord.query.one()
        assert sample.extra_metadata['unmatched_input'] is True and 'input_sample_id' not in sample.extra_metadata
        assert apply(client, project, response).json['rows'][0]['replayed'] is True
        source = client.get(f'/api/projects/{project.id}/input-samples?asset_set=甲').json
        assert '999' not in {row['sample_id'] for row in source['samples']}
        a = preview(client, project, [{'sample_id': '001', 'fields': {'illness': '甲'}}])
        b = preview(client, project, [{'sample_id': '001', 'fields': {'illness': '乙'}}], dataset='乙')
        apply(client, project, a)
        apply(client, project, b)
        assert {(row.extra_metadata['asset_set'], row.illness) for row in SampleRecord.query.filter_by(sample_id='001')} == {('甲', '甲'), ('乙', '乙')}


def test_duplicate_identifiers_need_concrete_record_ids(batch_context):
    app, project = batch_context
    first = record(project)
    second = record(project)
    with app.test_client() as client:
        response = preview(client, project, [{'sample_id': '001'}, {'sample_id': '001'}])
        assert response.json['counts'] == {'invalid': 2}
        response = preview(client, project, [{'sample_id': '001'}])
        assert response.json['counts'] == {'invalid': 1}
        response = preview(client, project, [{'record_id': first.id, 'sample_id': '001', 'fields': {'institution': '一'}},
                                            {'record_id': second.id, 'sample_id': '001', 'fields': {'institution': '二'}}])
        assert apply(client, project, response).json['counts'] == {'saved': 2}
        assert first.institution == '一' and second.institution == '二'


def test_record_id_cannot_move_across_dataset_or_identifier(batch_context):
    app, project = batch_context
    sample = record(project)
    with app.test_client() as client:
        for identifier, dataset in [('002', '甲'), ('001', '乙')]:
            response = preview(client, project, [{'record_id': sample.id, 'sample_id': identifier, 'fields': {'illness': '坏'}}], dataset)
            assert response.json['counts'] == {'invalid': 1}
            assert apply(client, project, response).json['counts'] == {'failed': 1}
        assert sample.illness is None


@pytest.mark.parametrize('row', [
    {'sample_id': '001', 'fields': {'sample_id': '002'}},
    {'sample_id': '001', 'clear_fields': ['样本名称']},
    {'sample_id': '001', 'clear_fields': ['疾病'], 'fields': {'illness': '冲突'}},
    {'sample_id': '001', 'fields': []},
    {'sample_id': '001', 'fields': {'illness': '字' * 256}},
    {'sample_id': '001', 'extra_metadata': {'asset_set': '乙'}},
    {'sample_id': 1},
])
def test_invalid_rows_never_mutate(batch_context, row):
    app, project = batch_context
    with app.test_client() as client:
        response = preview(client, project, [row])
        assert response.json['counts'] == {'invalid': 1}
        assert apply(client, project, response).json['counts'] == {'failed': 1}
        assert SampleRecord.query.count() == 0


def test_template_xlsx_preserves_text_ids_and_can_preview_round_trip(batch_context):
    app, project = batch_context
    sample = record(project, institution='机构')
    with app.test_client() as client:
        response = client.post(f'/api/projects/{project.id}/samples/batch/template', json={'asset_set': '甲'})
        assert response.status_code == 200
        workbook = load_workbook(io.BytesIO(response.data))
        sheet = workbook['样本登记']
        assert sheet['B2'].value == '001' and sheet['B2'].data_type == 's' and sheet['B2'].number_format == '@'
        assert sheet['B3'].value == '002'
        previewed = client.post(f'/api/projects/{project.id}/samples/batch/preview',
            data={'asset_set': '甲', 'file': (io.BytesIO(response.data), '模板.xlsx')})
        assert previewed.status_code == 200
        assert previewed.json['rows'][0]['status'] == 'unchanged'
        assert SampleRecord.query.count() == 1 and sample.institution == '机构'


def test_csv_and_numeric_excel_validation(batch_context):
    app, project = batch_context
    with app.test_client() as client:
        csv_content = '原始样本编号,样本名称,疾病\n001,登记,疾病甲\n'.encode('utf-8-sig')
        response = client.post(f'/api/projects/{project.id}/samples/batch/preview', data={'asset_set': '甲', 'file': (io.BytesIO(csv_content), '登记.csv')})
        assert response.json['rows'][0]['sample_id'] == '001'
        assert apply(client, project, response).json['counts'] == {'saved': 1}
        workbook = Workbook()
        workbook.active.append(['原始样本编号', '疾病'])
        workbook.active.append([1, '不覆盖'])
        buffer = io.BytesIO()
        workbook.save(buffer)
        response = client.post(f'/api/projects/{project.id}/samples/batch/preview', data={'asset_set': '甲', 'file': (io.BytesIO(buffer.getvalue()), '数字.xlsx')})
        assert response.json['rows'][0]['status'] == 'invalid'
        assert '文本' in response.json['rows'][0]['message']


def test_signed_preview_and_retry_rows_are_validated_before_any_write(batch_context):
    app, project = batch_context
    other = Project(name='其他项目')
    db.session.add(other)
    db.session.commit()
    with app.test_client() as client:
        response = preview(client, project, [{'sample_id': '001'}])
        assert apply(client, other, response).status_code == 400
        assert client.post(f'/api/projects/{project.id}/samples/batch/apply', json={'preview_token': response.json['preview_token'] + 'tampered'}).status_code == 400
        for rows in ([], [100], [1, 1], [True]):
            assert apply(client, project, response, rows=rows).status_code == 400
        assert SampleRecord.query.count() == 0


def test_input_change_after_preview_does_not_create_stale_registration(batch_context):
    app, project = batch_context
    with app.test_client() as client:
        response = preview(client, project, [{'sample_id': '001'}])
        asset = ProjectAsset.query.filter(ProjectAsset.metadata_json['asset_set'].as_string() == '甲').one()
        asset.metadata_json = {**asset.metadata_json, 'superseded': True}
        db.session.commit()
        result = apply(client, project, response)
        assert result.json['counts'] == {'failed': 1}
        assert SampleRecord.query.count() == 0


def test_normal_metadata_api_preserves_original_identifier(batch_context):
    app, project = batch_context
    sample = record(project)
    with app.test_client() as client:
        response = client.put(f'/api/samples/{sample.id}', json={'sample_id': '999'})
        assert response.status_code == 400 and sample.sample_id == '001'
        assert client.put(f'/api/samples/{sample.id}', json={'sample_id': '001', 'illness': '正常修改'}).status_code == 200
        assert sample.illness == '正常修改'


def test_account_bound_preview_and_template(batch_context):
    app, project = batch_context
    owner = User(username='owner', email='owner@example.test', password_hash='synthetic')
    other = User(username='other', email='other@example.test', password_hash='synthetic', role='admin')
    db.session.add_all([owner, other])
    db.session.flush()
    project.user_id = owner.id
    db.session.commit()
    app.config['REQUIRE_LOGIN'] = True
    with app.test_client() as client:
        with client.session_transaction() as session:
            session['_user_id'] = str(owner.id)
        response = preview(client, project, [{'sample_id': '001'}])
        assert response.status_code == 200
        with client.session_transaction() as session:
            session['_user_id'] = str(other.id)
        # The fixture retains one outer app context; clear its cached login user between sessions.
        g.pop('_login_user', None)
        assert apply(client, project, response).status_code == 400
        assert client.post(f'/api/projects/{project.id}/samples/batch/template', json={'asset_set': '甲'}).status_code == 400
        assert SampleRecord.query.count() == 0
        project.user_id = other.id
        db.session.commit()
        assert apply(client, project, response).status_code == 400  # Even a new owner cannot reuse another user's preview.

def test_selected_template_upload_only_applies_exact_selected_records(batch_context):
    import json
    app, project = batch_context
    selected = record(project, illness='所选旧值')
    unselected = record(project, identifier='002', illness='未选旧值')
    other_dataset = record(project, dataset='乙', illness='乙旧值')
    content = ('登记记录标识,原始样本编号,疾病\n'
               f'{selected.id},001,所选新值\n{unselected.id},002,不可修改\n'
               f'{other_dataset.id},001,不可跨集\n,003,不可新增\n').encode('utf-8-sig')
    with app.test_client() as client:
        response = client.post(f'/api/projects/{project.id}/samples/batch/preview', data={
            'asset_set': '甲', 'record_ids': json.dumps([selected.id]), 'allow_unmatched': 'true',
            'file': (io.BytesIO(content), '所选登记.csv')})
        assert response.status_code == 200
        assert response.json['counts'] == {'update': 1, 'invalid': 3}
        assert selected.illness == '所选旧值'
        assert apply(client, project, response).json['counts'] == {'saved': 1, 'failed': 3}
        assert selected.illness == '所选新值'
        assert unselected.illness == '未选旧值' and other_dataset.illness == '乙旧值'
        assert SampleRecord.query.count() == 3


@pytest.mark.parametrize('selection', [[], 'wrong', [False], ['missing'], ['x' * 37]])
def test_invalid_selected_scope_is_rejected_before_preview(batch_context, selection):
    app, project = batch_context
    sample = record(project, illness='保留')
    with app.test_client() as client:
        response = preview(client, project, [{'record_id': sample.id, 'sample_id': '001',
            'fields': {'illness': '不写入'}}], record_ids=selection)
        assert response.status_code == 400 and sample.illness == '保留'


def test_malformed_multipart_selection_is_rejected(batch_context):
    app, project = batch_context
    with app.test_client() as client:
        response = client.post(f'/api/projects/{project.id}/samples/batch/preview', data={
            'asset_set': '甲', 'record_ids': '[broken',
            'file': (io.BytesIO('原始样本编号,疾病\n001,不写入'.encode()), '登记.csv')})
        assert response.status_code == 400 and SampleRecord.query.count() == 0
